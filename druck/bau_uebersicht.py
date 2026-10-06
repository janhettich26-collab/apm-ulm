# Uebersicht-PDF "Was ist wofuer" + Word-Briefvorlage (nach bau_druck.py + druck.mjs)
import os, json, html
from PIL import Image, ImageOps
exec(open(os.path.expanduser('~/Websites/apm-ulm-v2/druck/bau_druck.py'), encoding='utf-8').read().split('# ============ 01 Logo')[0])
J = json.load(open(os.path.join(BAU, 'jobs.json')))

def vorschau(prefix, n=None, bleed_crop=True):
    """Vorschaubild(er) zu Dateien mit diesem Namensanfang, Beschnitt abgeschnitten."""
    out = []
    for j in J:
        if j['name'].startswith(prefix):
            for k, p in enumerate(j['prev']):
                im = Image.open(p).convert('RGB')
                if j['bleed']:
                    f = im.width / j['w']; c = round(j['bleed'] * f)
                    im = im.crop((c, c, im.width - c, im.height - c))
                q = os.path.join(BAU, f'ue_{len(out)}_{abs(hash(p))}.png'); im.save(q); out.append(q)
            if n and len(out) >= n: break
    return out[:n] if n else out

E = [  # (Titel, Ordner, Vorschau-Prefix, Anzahl Bilder, Format, Wofuer, Bestellen)
 ('Logo als Vektor', '01_Logo_Vektor', 'Logo_APM-Ulm_mit', 2, 'SVG und PDF, beliebig vergrößerbar',
  'Grunddateien mit „altun property management“ – in Navy, Weiß und Schwarz, auch ohne Unterzeile.',
  'An jede Druckerei, Folierer oder Stickerei schicken, wenn jemand „das Logo als Vektor“ will.'),
 ('Visitenkarte', '02_Visitenkarte', 'Visitenkarte', 2, '85 × 55 mm, Datei 91 × 61 mm (3 mm Beschnitt)',
  'Seite 1 = Vorderseite (Logo auf Navy), Seite 2 = Rückseite mit Name, Telefon, Mail, Web und Adresse.',
  'Online-Druckerei, 350 g Bilderdruck matt, beidseitig 4-farbig, 250–500 Stück.'),
 ('Flyer A5', '03_Flyer', 'Flyer_A5', 2, '148 × 210 mm, Datei 154 × 216 mm',
  'Vorne: Vorstellung und drei Vorteile. Hinten: alle 23 Leistungen, vier Schritte, Kontakt und QR-Code zur Webseite.',
  '135–170 g Bilderdruck matt, beidseitig, 500–1.000 Stück. Zum Mitgeben bei Besichtigungen.'),
 ('Flyer DIN lang', '03_Flyer', 'Flyer_DIN', 2, '99 × 210 mm, Datei 105 × 216 mm',
  'Schmales Format für Briefkästen, Hausverwaltungen und Auslagen. Hinten die Leistungsliste.',
  '135–170 g Bilderdruck, beidseitig, 1.000 Stück.'),
 ('Autotür – Logo', '04_Auto-Beschriftung', 'Aufkleber_Logo_Tuer', 2, '30, 45 oder 60 cm breit · Weiß oder Navy',
  'Logo für Fahrer- und Beifahrertür. Weiß für dunkle Autos, Navy für helle Autos.',
  'Folierer vor Ort: „Plottfolie, geschnitten, ohne Hintergrund“. Die SVG-Datei mitgeben. Die weiße Datei sieht am Bildschirm leer aus – das ist richtig.'),
 ('Transporter – Seitenbeschriftung', '04_Auto-Beschriftung', 'Aufkleber_Seitenbeschriftung', 2, '80 oder 100 cm breit · Weiß oder Navy',
  'Logo, Leistungen, Telefon und Webadresse für die Seitenfläche eines Transporters.',
  'Wie oben beim Folierer; vorher die freie Fläche am Auto ausmessen.'),
 ('Heckscheibe', '04_Auto-Beschriftung', 'Aufkleber_Heckscheibe', 1, '60, 80 oder 100 cm breit',
  'Eine Zeile: www.apm-ulm.de · 0170 5810174. Weiß wirkt auf getönten Scheiben am besten.',
  'Plottfolie. Unterkante der Scheibe nutzen, damit die Sicht frei bleibt.'),
 ('Magnetschild', '04_Auto-Beschriftung', 'Magnetschild', 2, '50 × 30 cm oder 60 × 40 cm (3 mm Beschnitt)',
  'Zum Abnehmen – ideal fürs private Auto. Logo, Leistungen, Telefon, Web auf Navy.',
  'Magnetfolie bedruckt, Ecken abgerundet, 2 Stück. Nur auf glattes Blech, vor der Waschanlage abnehmen.'),
 ('Stempel', '05_Stempel', 'Stempel', 3, '47 × 18 · 58 × 22 · 70 × 25 mm',
  'Firmenstempel schwarz mit Logo, Name, Adresse, Telefon, Mail (bei den größeren auch Web und USt-IdNr.).',
  'Empfehlung: Trodat Printy 4913 (58 × 22 mm). Datei im Stempelshop hochladen oder beim Schlüsseldienst abgeben.'),
 ('Briefpapier', '06_Briefpapier', 'Briefpapier_A4_zum', 1, 'A4 nach DIN 5008, mit Falz- und Lochmarken',
  'Logo oben rechts, Absenderzeile fürs Fensterkuvert, Fußzeile mit Kontakt und USt-IdNr. Dazu eine Word-Vorlage zum Schreiben.',
  'Druckerei: Datei „DRUCKEREI“, 90–100 g Offsetpapier. Oder einfach die Word-Vorlage nutzen und selbst drucken.'),
 ('Aushang Treppenhaus', '07_Aushang_Treppenhaus', 'Aushang', 1, 'A4',
  'Für das schwarze Brett in betreuten Häusern: Ansprechpartner, Telefon, Felder für Objekt und Rundgang zum Ausfüllen.',
  'Selbst ausdrucken, am besten laminieren.'),
 ('Aufkleber Technikraum', '08_Aufkleber_Technik', 'Aufkleber_Technikraum', 1, '70 × 35 mm',
  'Für Türen von Heizungs-, Technik- und Müllräumen: „Störung oder Schaden? Bitte melden“ mit Telefon und QR-Code.',
  'Etiketten auf weißer Folie, matt, wetterfest, 100 Stück.'),
 ('Aufkleber Wartungsnachweis', '08_Aufkleber_Technik', 'Aufkleber_Wartung', 1, '60 × 40 mm',
  'Zum Beschriften direkt an der Anlage: Anlage, letzte und nächste Wartung, Zeichen.',
  'Beschreibbare Folie (matt, mit Kugelschreiber beschreibbar), 100 Stück.'),
 ('Arbeitskleidung', '09_Arbeitskleidung', 'Textil', 3, 'Brust 9 cm · Rücken 28 cm',
  'Brust links mit Unterzeile für Druck, ohne Unterzeile für Stick (die feine Schrift lässt sich nicht sticken). Rücken mit Webadresse.',
  'Textildruck oder Stickerei vor Ort. Weiß für Navy- oder dunkle Kleidung, Navy für helle Kleidung.'),
 ('Banner Bauzaun / Gerüst', '10_Banner_Bauzaun', 'Banner', 1, '200 × 100 cm',
  'Für Baustellen und Modernisierungen: Logo, Leistungen, Telefon, Web – aus 30 m Entfernung lesbar.',
  'PVC-Banner 510 g, Ösen rundum alle 50 cm. Datei ist Endformat; der 5-cm-Rand ist für Ösen frei.'),
]

def bild_html(ps):
    hoch = all(Image.open(p).height > Image.open(p).width for p in ps)
    cls = 'hoch' if hoch else f'n{len(ps)}'
    return f'<div class="bi {cls}">' + ''.join(f'<img src="file://{p}">' for p in ps) + '</div>'

reihen = []
for t, ordner, pre, n, fmt, wofuer, best in E:
    reihen.append(f'''<div class="it">{bild_html(vorschau(pre, n))}
<div class="tx"><h3>{t}</h3><div class="or">Ordner {ordner}</div>
<p><b>Format</b> {fmt}</p><p><b>Wofür</b> {wofuer}</p><p><b>Bestellen</b> {best}</p></div></div>''')

ll, _ = logo(NAVY, True)
seite1 = f'''<div class="kopf"><div style="width:52mm">{ll}</div><div class="dt">Druckvorlagen · Stand 06.10.2026</div></div>
<h1>Alle Druckvorlagen <em>auf einen Blick</em></h1>
<p class="lead">Jede Datei heißt so, wie sie benutzt wird. Dateien mit <b>DRUCK</b> im Namen gehen direkt an die Druckerei, Dateien mit <b>Folie</b> an den Folierer, <b>SVG</b> ist die Vektorfassung für Folie, Stick und Stempel.</p>
<div class="box"><h4>Das gilt für alle Dateien</h4>
<p><b>Farbe:</b> Navy #1C2840 (RGB 28 / 40 / 64). Für Folie und Textil nach einem Farbton nahe <b>Pantone 533 C</b> fragen und Muster ansehen.</p>
<p><b>Schrift:</b> Verdana. Das Logo ist reine Vektorgrafik; bei Folie, Stempel und Textil ist auch die Schrift in Kurven umgewandelt.</p>
<p><b>Beschnitt:</b> Druckdateien haben rundum 3 mm mehr Hintergrund, der beim Schneiden wegfällt. Bei der Online-Druckerei einfach das Endformat wählen (z. B. 85 × 55 mm) und die Datei hochladen.</p>
<p><b>Farbraum:</b> Die PDFs sind in RGB. Online-Druckereien wandeln das beim Hochladen automatisch in Druckfarben um; das Navy kann gedruckt minimal dunkler wirken.</p>
<p><b>E-Mail-Signatur:</b> kein Druck – liegt fertig unter www.apm-ulm.de/signatur.html.</p></div>
<table><tr><th>Ordner</th><th>Inhalt</th></tr>
{''.join(f'<tr><td>{o}</td><td>{t}</td></tr>' for o, t in [
 ('01_Logo_Vektor','Logo mit und ohne Unterzeile, Navy / Weiß / Schwarz'),
 ('02_Visitenkarte','Visitenkarte Emre Altun, beidseitig'),
 ('03_Flyer','Flyer A5 und DIN lang, beidseitig'),
 ('04_Auto-Beschriftung','Tür-Logo, Transporter-Seite, Heckscheibe, Magnetschild'),
 ('05_Stempel','Drei Stempelgrößen, schwarz'),
 ('06_Briefpapier','Briefbogen für Druckerei, zum Selbstdrucken und als Word-Vorlage'),
 ('07_Aushang_Treppenhaus','Ansprechpartner-Aushang A4'),
 ('08_Aufkleber_Technik','Technikraum-Aufkleber, Wartungsnachweis'),
 ('09_Arbeitskleidung','Brust- und Rückenlogo für Druck und Stick'),
 ('10_Banner_Bauzaun','PVC-Banner 200 × 100 cm')])}</table>'''
pro_seite = 4
seiten_ = [seite1] + [''.join(reihen[i:i + pro_seite]) for i in range(0, len(reihen), pro_seite)]
css = f'''.seite{{padding:16mm 16mm 14mm}} .kopf{{display:flex;justify-content:space-between;align-items:flex-start}} .dt{{font-size:7.5pt;color:{GRAU};margin-top:2mm}}
h1{{font-weight:400;font-size:22pt;margin-top:14mm;line-height:1.2}} h1 em{{font-style:normal;color:{GRAU}}}
.lead{{font-size:9.5pt;line-height:1.6;color:{GRAU};margin-top:5mm}} .lead b{{color:{NAVY}}}
.box{{margin-top:8mm;padding:5mm 6mm;border-left:.8mm solid {NAVY};background:#F3F4F7}} .box h4{{font-size:9pt;margin-bottom:2.5mm}}
.box p{{font-size:8pt;line-height:1.55;margin-top:1.6mm;color:{GRAU}}} .box b{{color:{NAVY}}}
table{{margin-top:8mm;width:100%;border-collapse:collapse;font-size:8pt}} th{{text-align:left;font-size:7pt;letter-spacing:.08em;color:{GRAU};font-weight:400;padding:0 0 2mm}}
td{{padding:2mm 0;border-top:.2mm solid {LINIE}}} td:first-child{{width:58mm;font-weight:700}}
.it{{display:flex;gap:7mm;padding:5mm 0;border-bottom:.2mm solid {LINIE};height:64mm}} .it:last-child{{border:0}}
.bi{{width:72mm;flex:none;display:flex;flex-wrap:wrap;gap:2mm;align-content:flex-start}}
.bi img{{max-width:72mm;max-height:26mm;border:.2mm solid {LINIE}}} .bi img:only-child{{max-height:54mm}} .bi.n3 img{{max-height:16.5mm}} .bi.hoch img{{max-width:34mm;max-height:54mm}}
.tx h3{{font-size:11pt}} .or{{font-size:7pt;color:{GRAU};margin-top:1mm;letter-spacing:.04em}}
.tx p{{font-size:8pt;line-height:1.5;margin-top:2mm;color:{GRAU}}} .tx b{{color:{NAVY};display:inline-block;width:19mm}}'''
JOBS.clear()
seiten('00_Uebersicht_Was-ist-wofuer.pdf', 210, 297, seiten_, css)
json.dump(JOBS, open(os.path.join(BAU, 'jobs_ue.json'), 'w'), ensure_ascii=False)

# ---------- Word-Briefvorlage ----------
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
doc = Document(); sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Mm(25), Mm(20), Mm(45), Mm(30)
sec.header_distance, sec.footer_distance = Mm(12), Mm(10)
st = doc.styles['Normal']; st.font.name = 'Verdana'; st.font.size = Pt(10); st.font.color.rgb = RGBColor(0x1C, 0x28, 0x40)
st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.25
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.add_run().add_picture(os.path.join(BAU, 'logo_brief.png'), width=Mm(56))
def absatz(t='', size=None, bold=False, farbe=None, rechts=False, vor=0):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = bold
    if size: r.font.size = Pt(size)
    if farbe: r.font.color.rgb = RGBColor.from_string(farbe)
    if rechts: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(vor); return p
absatz('apm ulm · altun property management · Weinbergweg 81 · 89075 Ulm', 6.5, farbe='566078')
for z in ['Firma', 'Herr / Frau Name', 'Straße Hausnummer', 'PLZ Ort']: absatz(z, vor=0)
absatz('Ulm, TT.MM.JJJJ', rechts=True, vor=40)
absatz('Betreff', bold=True, vor=24)
absatz('Sehr geehrte Damen und Herren,', vor=18)
absatz('hier steht Ihr Text.', vor=12)
absatz('Mit freundlichen Grüßen', vor=18)
absatz(NAME, vor=30)
absatz(f'{ROLLE} · {FIRMA}', 8, farbe='566078')
ft = sec.footer.add_table(rows=1, cols=3, width=Mm(165))
for c, t in zip(ft.rows[0].cells, [f'{FIRMA}\nInhaber {NAME}\n{STR} · {ORT}', f'Telefon {TEL_INT}\n{INFO}\n{WEB}', f'USt-IdNr. {UST}']):
    c.text = ''; r = c.paragraphs[0].add_run(t); r.font.size = Pt(6.5); r.font.color.rgb = RGBColor(0x56, 0x60, 0x78)
os.makedirs(os.path.join(OUT, '06_Briefpapier'), exist_ok=True)
doc.save(os.path.join(OUT, '06_Briefpapier', 'Briefvorlage_Word_zum-Schreiben.docx'))
print('Übersicht + Word fertig')
